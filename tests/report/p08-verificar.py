"""Leitura independente do XLSX do app; não cria nem altera a planilha."""
import json,re
from pathlib import Path
from openpyxl import load_workbook
w=load_workbook('output/xlsx/ninhos-p08-demonstracao.xlsx',data_only=False)
s=w['Ninhos'];esperado=json.loads(Path('tmp/xlsx/esperado.json').read_text(encoding='utf-8'))
assert s.freeze_panes=='B2'
assert [s.cell(1,c).value for c in range(2,7)]==['001','002','003','004','005']
assert all(c.data_type!='f' for sh in w for row in sh for c in row)
falhas=[];quant=0
for col,n in enumerate(esperado,2):
 valores={};ultimo=None
 for row in s.iter_rows(min_row=2):
  label=row[0].value;v=row[col-1].value
  if row[0].style_id==2:ultimo=None;continue
  if label and re.search(r' \(continuação \d+\)$',label) and ultimo:
   valores[ultimo][-1]+=str(v or '')
  elif label:
   ultimo=label;valores.setdefault(label,[]).append(str(v) if v is not None else '')
 for label,v in n['linhas']:
  quant+=1;opcoes=valores.get(label,[])
  if v in opcoes:continue
  try:
   if any(abs(float(v)-float(x))<1e-9 for x in opcoes if re.fullmatch(r'-?\d+(\.\d+)?',x)):continue
  except ValueError:pass
  falhas.append((n['numero'],label,v,opcoes))
assert not falhas,repr(falhas[:8])
assert any(c.data_type=='n' and c.value==0 for row in s for c in row)
print(f'{quant} rótulos/valores presentes; 5 ninhos em colunas; zero numérico; sem fórmulas; congelamento B2; {s.max_row} linhas.')
